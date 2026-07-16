#!/usr/bin/env python3
"""Convert episode script-draft.md files to .docx using the channel's
RAMP_[Topic]_VideoScript.docx naming convention.

Usage:
    python3 tools/scripts_to_docx.py            # convert all restaurant-track scripts
    python3 tools/scripts_to_docx.py <path.md>  # convert one script
Output lands next to each source script.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

REPO = Path(__file__).resolve().parent.parent

DOCX_NAMES = {
    "video-01-taking-orders": "RAMP_TakingOrders_VideoScript.docx",
    "video-02-angry-customer": "RAMP_AngryCustomer_VideoScript.docx",
    "video-03-kitchen-english": "RAMP_KitchenEnglish_VideoScript.docx",
    "video-04-calling-in-sick": "RAMP_CallingInSick_VideoScript.docx",
    "video-05-shift-handover": "RAMP_ShiftHandover_VideoScript.docx",
    "video-06-blamed-for-kitchen": "RAMP_BlamedForKitchen_VideoScript.docx",
    "video-07-english-freezes": "RAMP_EnglishFreezes_VideoScript.docx",
}


def add_runs(paragraph, text):
    """Split **bold** and *italic* markdown into styled runs."""
    for part in re.split(r"(\*\*.*?\*\*|\*.*?\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def convert(md_path: Path, out_path: Path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for raw in md_path.read_text().splitlines():
        line = raw.rstrip()
        if line.strip() == "---":
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(re.sub(r"\*", "", m.group(2)), level=len(m.group(1)))
            continue
        if re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^\s*[-*]\s+", "", line))
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\s*\d+\.\s+", "", line))
            continue
        if line.startswith("|"):  # simple table row -> plain text line
            p = doc.add_paragraph()
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.match(r"^-+$", c) for c in cells if c):
                continue
            add_runs(p, "   ·   ".join(c for c in cells if c))
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            run_text = line[2:]
            add_runs(p, run_text)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x88, 0x55, 0x00)
            continue
        p = doc.add_paragraph()
        add_runs(p, line)

    doc.save(out_path)
    print(f"wrote {out_path}")


def main():
    if len(sys.argv) > 1:
        targets = [Path(sys.argv[1])]
    else:
        targets = sorted(REPO.glob("restaurant-track/video-*/script-draft.md"))
    for md in targets:
        folder = md.parent.name
        name = DOCX_NAMES.get(folder, f"RAMP_{folder}_VideoScript.docx")
        convert(md, md.parent / name)


if __name__ == "__main__":
    main()
